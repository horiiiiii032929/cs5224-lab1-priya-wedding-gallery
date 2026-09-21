from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "Lab1.docx"
OUTPUT = ROOT / "Lab1-editable.docx"
IMAGE = ROOT / "diagram" / "priya-architecture-drawio.png"
REPOSITORY_URL = "https://github.com/horiiiiii032929/cs5224-lab1-priya-wedding-gallery"


def image_target(document: Document, paragraph) -> str | None:
    for blip in paragraph._p.iter(qn("a:blip")):
        rel_id = blip.get(qn("r:embed"))
        if rel_id and rel_id in document.part.rels:
            return document.part.rels[rel_id].target_ref
    return None


def remove_paragraph(paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


doc = Document(SOURCE)


def all_paragraphs(document: Document):
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


# Keep Part 2 explicitly framed as a proposed design and tighten claims that
# could otherwise imply the complete architecture is already implemented in
# the repository. Replacements are made within runs to preserve formatting.
replacements = {
    "Priya's system is a private S3 bucket behind Cognito sign-in. The only custom code is one small \"add couple\" function; every other piece is a managed AWS service.":
        "Part 2 proposes a private S3 bucket behind Cognito sign-in. The only custom backend business logic would be the small \"add couple\" function; the remaining components are managed AWS services. The full system was designed, but not deployed, for this lab.",
    "the backend (Cognito, S3, Lambda) is defined as TypeScript code and deployed with AWS CDK":
        "the proposed backend (Cognito, S3, Lambda) would be defined as TypeScript code and deployed with AWS CDK",
    "Cognito's hosted pages": "Cognito managed login pages",
    "AWS hosts the sign-in pages; couples sign in with a one-time email code.":
        "Cognito managed login handles sign-in; couples use a one-time email code. The user pool uses the Essentials tier, with email delivery configured through Amazon SES.",
    "Gen 2 also defines Cognito, S3 and Lambda as TypeScript code (defineAuth, defineStorage, defineFunction) deployed by CDK, so the whole system can be rebuilt from Git.":
        "In an implementation, Gen 2 would define Cognito, S3 and Lambda as TypeScript code (defineAuth, defineStorage, defineFunction) deployed by CDK, so the configuration could be version-controlled in Git.",
    "The only custom code, about 20 lines.":
        "The only custom backend business logic, about 20 lines.",
    "It's free up to 10,000 monthly active users; Priya has about 80 couples a year.":
        "Cognito is free up to 10,000 monthly active users; SES email charges are separate but negligible at this scale. Priya has about 80 couples a year.",
    " Lambda. Couples sign in":
        " Lambda. The function accepts only lowercase letters, digits and hyphens for wedding_id, preventing IAM wildcard characters from entering the principal tag. Couples sign in",
    "The identity pool issues credentials that expire after 1 hour, and Storage Browser uses them to sign each download, so any download link copied out of the app stops working within the hour.":
        "The identity pool issues temporary credentials, and Storage Browser signs S3 requests or generates short-lived presigned URLs. A copied URL expires at its configured expiry or sooner when the underlying role session expires.",
}

for paragraph in all_paragraphs(doc):
    for run in paragraph.runs:
        for old, new in replacements.items():
            if old in run.text:
                run.text = run.text.replace(old, new)

# Remove the redundant caption above the landscape figure. The figure itself
# now contains the requested title and region subtitle.
for paragraph in list(doc.paragraphs):
    if paragraph.text.startswith("Figure 1. Architecture of Priya"):
        remove_paragraph(paragraph)
        break

# Replace only the architecture image, preserving every other report element.
for paragraph in doc.paragraphs:
    if image_target(doc, paragraph) == "media/image7.png":
        paragraph.clear()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = 0
        paragraph.paragraph_format.space_after = 0
        run = paragraph.add_run()
        run.add_picture(str(IMAGE), width=Inches(10.68), height=Inches(7.18))
        break
else:
    raise RuntimeError("Architecture image was not found in the source report")

# Add the source repository URL at the bottom without changing existing prose
# or code snippets.
repository_paragraph = doc.add_paragraph()
repository_paragraph.paragraph_format.space_before = 8
repository_run = repository_paragraph.add_run(
    f"Repository (Part 1 proof-of-concept and editable report assets): {REPOSITORY_URL}"
)
repository_run.font.size = Pt(9)

ai_paragraph = doc.add_paragraph()
ai_paragraph.paragraph_format.space_before = 4
ai_run = ai_paragraph.add_run(
    "AI declaration: Claude Opus 4.5 and OpenAI GPT-5.6 Sol "
    "(medium reasoning) were used to assist with drafting, formatting, and "
    "diagram refinement."
)
ai_run.font.size = Pt(9)

# Keep the existing three-section A4 structure: portrait, landscape, portrait.
doc.save(OUTPUT)
print(OUTPUT)
